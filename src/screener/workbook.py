from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.screener.engine import ScreenerEngine

OUTPUT_PATH = Path("output/screener_output.xlsx")


def build_workbook():
    """Build workbook."""
    engine = ScreenerEngine()

    df, results = engine.run_all()

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    preset_order = [
        "quality_compounder",
        "value_pick",
        "growth_accelerator",
        "dividend_champion",
        "debt_free_blue_chip",
        "turnaround_watch",
    ]

    sheet_names = {
        "quality_compounder": "Quality Compounder",
        "value_pick": "Value Pick",
        "growth_accelerator": "Growth Accelerator",
        "dividend_champion": "Dividend Champion",
        "debt_free_blue_chip": "Debt-Free Blue Chip",
        "turnaround_watch": "Turnaround Watch",
    }

    # Required KPI columns for Sprint 3.
    columns = [
        "company_id",
        "company_name",
        "sector",
        "industry",
        "year",
        "market_cap_crore",
        "pe_ratio",
        "pb_ratio",
        "dividend_yield_pct",
        "return_on_equity_pct",
        "return_on_capital_employed_pct",
        "net_profit_margin_pct",
        "operating_profit_margin_pct",
        "debt_to_equity",
        "interest_coverage",
        "asset_turnover",
        "free_cash_flow_cr",
        "cash_from_operations_cr",
        "cfo_pat_ratio",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "eps_cagr_5yr",
        "earnings_per_share",
        "book_value_per_share",
        "dividend_payout_ratio_pct",
        "profitability_score",
        "cash_quality_score",
        "growth_score",
        "leverage_score",
        "sprint3_composite_score",
    ]

    available_columns = [col for col in columns if col in df.columns]

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:

        # Summary sheet.
        summary_rows = []

        for key in preset_order:
            result = results[key].copy()

            summary_rows.append(
                {
                    "preset": sheet_names[key],
                    "companies": len(result),
                    "average_score": (
                        result["sprint3_composite_score"].mean()
                        if not result.empty
                        else None
                    ),
                    "highest_score": (
                        result["sprint3_composite_score"].max()
                        if not result.empty
                        else None
                    ),
                }
            )

        summary = pd.DataFrame(summary_rows)

        summary.to_excel(writer, sheet_name="Summary", index=False)

        # Six preset sheets.
        for key in preset_order:

            result = results[key].copy()

            if not result.empty:
                result = result.sort_values("sprint3_composite_score", ascending=False)

            result = result[available_columns]

            result.to_excel(writer, sheet_name=sheet_names[key], index=False)

    # Excel formatting.
    wb = load_workbook(OUTPUT_PATH)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")

    header_font = Font(color="FFFFFF", bold=True)

    green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")

    red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

    yellow_fill = PatternFill(fill_type="solid", fgColor="FFEB9C")

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in ws.columns:

            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 30)

        # Composite score highlighting.
        if ws.title != "Summary":
            headers = {cell.value: cell.column for cell in ws[1]}

            if "sprint3_composite_score" in headers:

                score_col = headers["sprint3_composite_score"]

                for row in range(2, ws.max_row + 1):

                    cell = ws.cell(row=row, column=score_col)

                    if isinstance(cell.value, (int, float)):

                        if cell.value >= 75:
                            cell.fill = green_fill

                        elif cell.value < 40:
                            cell.fill = red_fill

                        else:
                            cell.fill = yellow_fill

            # Percentage / score number formatting.
            for header, column in headers.items():

                for row in range(2, ws.max_row + 1):

                    cell = ws.cell(row=row, column=column)

                    if header.endswith(("_pct", "_score")) or header in [
                        "pe_ratio",
                        "pb_ratio",
                        "asset_turnover",
                        "debt_to_equity",
                        "interest_coverage",
                    ]:
                        cell.number_format = "0.00"

                    elif header in [
                        "market_cap_crore",
                        "free_cash_flow_cr",
                        "cash_from_operations_cr",
                    ]:
                        cell.number_format = "#,##0.00"

    # Summary formatting.
    ws = wb["Summary"]

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"

    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 15), 30)

    for row in range(2, ws.max_row + 1):

        companies = ws.cell(row=row, column=2).value
        avg_score = ws.cell(row=row, column=3).value

        if isinstance(companies, (int, float)):
            ws.cell(row=row, column=2).number_format = "0"

        if isinstance(avg_score, (int, float)):
            ws.cell(row=row, column=3).number_format = "0.00"

    wb.save(OUTPUT_PATH)

    print("=" * 70)
    print("SPRINT 3 SCREENER WORKBOOK")
    print("=" * 70)
    print(f"Created: {OUTPUT_PATH}")
    print()

    for key in preset_order:
        print(f"{sheet_names[key]}: " f"{len(results[key])} companies")

    print()
    print("Sheets:")
    for ws in wb.sheetnames:
        print(f"  - {ws}")

    print("=" * 70)


if __name__ == "__main__":
    build_workbook()

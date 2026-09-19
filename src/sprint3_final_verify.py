from pathlib import Path
import sys
import sqlite3
import pandas as pd
from openpyxl import load_workbook

# ------------------------------------------------------------------
# PROJECT ROOT / IMPORT FIX
# ------------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DB_PATH = PROJECT_ROOT / "data" / "nifty100.db"
SCREENER_PATH = PROJECT_ROOT / "output" / "screener_output.xlsx"
PEER_PATH = PROJECT_ROOT / "output" / "peer_comparison.xlsx"
RADAR_DIR = PROJECT_ROOT / "reports" / "radar_charts"

conn = sqlite3.connect(DB_PATH)

checks = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    checks.append((name, condition))
    print(f"[{status}] {name}")
    if detail:
        print(f"       {detail}")


print("=" * 75)
print("SPRINT 3 — DAY 21 FINAL VERIFICATION")
print("=" * 75)
print()


# ------------------------------------------------------------------
# DQ-01: peer_percentiles table exists
# ------------------------------------------------------------------

tables = {
    row[0]
    for row in conn.execute(
        """
        SELECT name
        FROM sqlite_master
        WHERE type='table'
        """
    ).fetchall()
}

check(
    "DQ-01 peer_percentiles table exists",
    "peer_percentiles" in tables
)


# ------------------------------------------------------------------
# DQ-02: percentile row count
# ------------------------------------------------------------------

peer_rows = conn.execute(
    "SELECT COUNT(*) FROM peer_percentiles"
).fetchone()[0]

check(
    "DQ-02 percentile row count",
    peer_rows == 560,
    f"Expected 560, found {peer_rows}"
)


# ------------------------------------------------------------------
# DQ-03: peer company count
# ------------------------------------------------------------------

peer_companies = conn.execute(
    "SELECT COUNT(DISTINCT company_id) FROM peer_percentiles"
).fetchone()[0]

check(
    "DQ-03 peer company count",
    peer_companies == 56,
    f"Expected 56, found {peer_companies}"
)


# ------------------------------------------------------------------
# DQ-04: peer group count
# ------------------------------------------------------------------

peer_groups = conn.execute(
    "SELECT COUNT(DISTINCT peer_group) FROM peer_percentiles"
).fetchone()[0]

check(
    "DQ-04 peer group count",
    peer_groups == 11,
    f"Expected 11, found {peer_groups}"
)


# ------------------------------------------------------------------
# DQ-05: metric count
# ------------------------------------------------------------------

metrics = conn.execute(
    "SELECT COUNT(DISTINCT metric) FROM peer_percentiles"
).fetchone()[0]

check(
    "DQ-05 metric count",
    metrics == 10,
    f"Expected 10, found {metrics}"
)


# ------------------------------------------------------------------
# DQ-06: percentile range
# ------------------------------------------------------------------

min_pct, max_pct = conn.execute(
    """
    SELECT MIN(percentile_rank), MAX(percentile_rank)
    FROM peer_percentiles
    """
).fetchone()

range_valid = (
    min_pct is not None
    and max_pct is not None
    and min_pct >= 0
    and max_pct <= 1
)

check(
    "DQ-06 percentile range",
    range_valid,
    f"Range = {min_pct} to {max_pct}"
)


# ------------------------------------------------------------------
# DQ-07: duplicate percentile records
# ------------------------------------------------------------------

duplicate_groups = conn.execute(
    """
    SELECT COUNT(*)
    FROM (
        SELECT company_id, peer_group, metric, year, COUNT(*) AS cnt
        FROM peer_percentiles
        GROUP BY company_id, peer_group, metric, year
        HAVING COUNT(*) > 1
    )
    """
).fetchone()[0]

check(
    "DQ-07 duplicate percentile records",
    duplicate_groups == 0,
    f"Duplicate groups = {duplicate_groups}"
)


# ------------------------------------------------------------------
# DQ-08: 10 metrics per company
# ------------------------------------------------------------------

metric_failures = conn.execute(
    """
    SELECT COUNT(*)
    FROM (
        SELECT company_id, COUNT(DISTINCT metric) AS metric_count
        FROM peer_percentiles
        GROUP BY company_id
        HAVING metric_count != 10
    )
    """
).fetchone()[0]

check(
    "DQ-08 10 metrics per company",
    metric_failures == 0,
    f"Companies failing = {metric_failures}"
)


# ------------------------------------------------------------------
# DQ-09: peer-group assignments populated
# ------------------------------------------------------------------

empty_groups = conn.execute(
    """
    SELECT COUNT(*)
    FROM peer_groups
    WHERE peer_group IS NULL
       OR TRIM(peer_group) = ''
    """
).fetchone()[0]

check(
    "DQ-09 peer-group assignments populated",
    empty_groups == 0,
    f"Empty groups = {empty_groups}"
)


# ------------------------------------------------------------------
# DQ-10: screener workbook exists
# ------------------------------------------------------------------

check(
    "DQ-10 screener_output.xlsx exists",
    SCREENER_PATH.exists(),
    str(SCREENER_PATH.relative_to(PROJECT_ROOT))
)


# ------------------------------------------------------------------
# DQ-11: six screener preset sheets
# ------------------------------------------------------------------

expected_screener_sheets = {
    "Summary",
    "Quality Compounder",
    "Value Pick",
    "Growth Accelerator",
    "Dividend Champion",
    "Debt-Free Blue Chip",
    "Turnaround Watch",
}

if SCREENER_PATH.exists():
    wb = load_workbook(SCREENER_PATH, read_only=True)
    actual_screener_sheets = set(wb.sheetnames)
    wb.close()

    missing_screener = sorted(
        expected_screener_sheets - actual_screener_sheets
    )

    check(
        "DQ-11 six screener preset sheets",
        not missing_screener,
        f"Missing = {missing_screener}"
    )
else:
    check(
        "DQ-11 six screener preset sheets",
        False,
        "Workbook missing"
    )


# ------------------------------------------------------------------
# DQ-12: eleven peer workbook sheets
# ------------------------------------------------------------------

expected_peer_sheets = {
    "Automobiles",
    "Consumer Finance",
    "FMCG",
    "IT Services",
    "Life Insurance",
    "Oil & Gas",
    "Pharmaceuticals",
    "Power & Utilities",
    "Private Banks",
    "Public Sector Banks",
    "Steel",
}

if PEER_PATH.exists():
    wb = load_workbook(PEER_PATH, read_only=True)
    actual_peer_sheets = set(wb.sheetnames)
    wb.close()

    missing_peer = sorted(
        expected_peer_sheets - actual_peer_sheets
    )

    check(
        "DQ-12 eleven peer workbook sheets",
        not missing_peer,
        f"Missing = {missing_peer}"
    )
else:
    check(
        "DQ-12 eleven peer workbook sheets",
        False,
        "Workbook missing"
    )


# ------------------------------------------------------------------
# DQ-13: radar chart count
# ------------------------------------------------------------------

if RADAR_DIR.exists():
    radar_count = len(
        list(RADAR_DIR.glob("*.png"))
    )
else:
    radar_count = 0

check(
    "DQ-13 radar chart count",
    radar_count == 56,
    f"Expected 56, found {radar_count}"
)


# ------------------------------------------------------------------
# DQ-14: database integrity and foreign keys
# ------------------------------------------------------------------

integrity = conn.execute(
    "PRAGMA integrity_check"
).fetchone()[0]

try:
    foreign_key_errors = conn.execute(
        "PRAGMA foreign_key_check"
    ).fetchall()
except sqlite3.DatabaseError:
    foreign_key_errors = [("foreign_key_check_error",)]

check(
    "DQ-14 database integrity",
    integrity == "ok" and len(foreign_key_errors) == 0,
    f"integrity={integrity}, FK failures={len(foreign_key_errors)}"
)


# ------------------------------------------------------------------
# SPRINT 3 ACCEPTANCE CHECKS
# ------------------------------------------------------------------

print()
print("=" * 75)
print("SPRINT 3 ACCEPTANCE CHECKS")
print("=" * 75)
print()


# ------------------------------------------------------------------
# Acceptance 1:
# Quality Compounder top 5 must have ROE > 15 and
# non-Financial companies must have D/E < 1
# ------------------------------------------------------------------

try:
    from src.screener.engine import ScreenerEngine

    engine = ScreenerEngine()
    df, results = engine.run_all()

    quality = results["quality_compounder"].copy()

    top5 = quality.sort_values(
        "sprint3_composite_score",
        ascending=False
    ).head(5)

    quality_valid = (
        len(top5) == 5
        and (top5["return_on_equity_pct"] > 15).all()
    )

    non_financial = ~(
        top5["sector"]
        .astype(str)
        .str.strip()
        .eq("Financials")
    )

    debt_valid = (
        (top5.loc[non_financial, "debt_to_equity"] < 1)
        .all()
    )

    check(
        "Acceptance-01 Quality Compounder top 5",
        quality_valid and debt_valid,
        f"Top 5 companies = {top5['company_id'].tolist()}"
    )

except Exception as exc:

    check(
        "Acceptance-01 Quality Compounder top 5",
        False,
        str(exc)
    )


# ------------------------------------------------------------------
# Acceptance 2:
# IT Services highest ROE gets highest ROE percentile
# ------------------------------------------------------------------

it_data = pd.read_sql_query(
    """
    SELECT
        pp.company_id,
        pp.value,
        pp.percentile_rank
    FROM peer_percentiles pp
    WHERE pp.peer_group = 'IT Services'
      AND pp.metric = 'roe'
    """,
    conn
)

if not it_data.empty:

    max_roe = it_data["value"].max()

    highest_roe = it_data[
        it_data["value"] == max_roe
    ]

    highest_percentile = (
        it_data["percentile_rank"].max()
    )

    it_valid = (
        (highest_roe["percentile_rank"] == highest_percentile)
        .all()
    )

    check(
        "Acceptance-02 IT Services ROE percentile",
        it_valid,
        f"Highest ROE={max_roe}, highest percentile={highest_percentile}"
    )

else:

    check(
        "Acceptance-02 IT Services ROE percentile",
        False,
        "No IT Services ROE data found"
    )


# ------------------------------------------------------------------
# FINAL RESULT
# ------------------------------------------------------------------

conn.close()

passed = sum(
    1 for _, condition in checks
    if condition
)

total = len(checks)

print()
print("=" * 75)
print(f"FINAL RESULT: {passed}/{total} CHECKS PASSED")
print("=" * 75)

if passed == total:
    print()
    print("SPRINT 3 VERIFICATION: SUCCESS")
    print()
    raise SystemExit(0)

else:
    print()
    print("SPRINT 3 VERIFICATION: FAIL")
    print()
    raise SystemExit(1)
